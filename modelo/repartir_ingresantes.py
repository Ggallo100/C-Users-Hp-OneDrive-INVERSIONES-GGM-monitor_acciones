#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Reparte un total de ingresantes proyectado por semestre, sede, carrera y
modalidad entre los CICLOS en que se producen realmente las convalidaciones.

    python3 repartir_ingresantes.py entrada.xlsx [salida.xlsx]
    python3 repartir_ingresantes.py --tabla        # sólo el cuadro de tasas

El archivo de entrada necesita las columnas Periodo, Sede, Carrera, Modalidad
y Nuevos. NO lleva columna Ciclo: eso es lo que añade este programa. El
resultado tiene el formato exacto que espera el modelo de proyección, de modo
que se puede cargar directamente en `modelo_proyeccion_matricula.html`.

El reparto se estima en dos partes, porque el histórico dice cosas distintas
de cada una:

  NIVEL   cuántos entran por encima del ciclo 1. Va del 0,2 % de Obstetricia
          al 19,7 % de Contabilidad, y es más del doble en las modalidades no
          presenciales. Depende además de la paridad del semestre.
  FORMA   a qué ciclo llegan los que entran alto. Es casi universal: 57 % al
          ciclo 2, 21 % al 3, 13 % al 4 y 9 % al 5 o más.

Ambas se contraen jerárquicamente hacia niveles más agregados, de modo que una
combinación con pocos ingresantes históricos hereda el comportamiento de su
carrera o de su modalidad en vez de inventarse una tasa propia.

El reparto entero usa el método del resto mayor: la suma por fila coincide
exactamente con el total declarado, sin perder ni inventar estudiantes.
"""
import json
import os
import sys

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import estimar as E

AQUI = os.path.dirname(os.path.abspath(__file__))
PAR = os.path.join(AQUI, "parametros.json")
CAB = ["Periodo", "Sede", "Carrera", "Modalidad", "Ciclo", "Nuevos", "CiclosPlan"]
CABF = PatternFill("solid", fgColor="D9E2F3")


def rotulo(p):
    return "%d-%s" % (p // 100, "I" if p % 100 == 1 else "II")


def leer_periodo(v):
    """Admite 2027-I, 2027-II, 202701 y 202702."""
    t = str(v).strip().upper().replace(" ", "")
    if "-" in t:
        a, r = t.split("-", 1)
        return int(a) * 100 + (2 if r.startswith("II") else 1)
    n = int(float(t))
    return n if n > 9999 else n


def resto_mayor(total, pesos):
    """
    Reparte un entero entre celdas conservando la suma exacta.

    Asignar `round(total * peso)` celda a celda no suma el total: el método del
    resto mayor da la parte entera a todos y reparte las unidades que faltan
    entre las celdas de mayor parte decimal.
    """
    total = int(round(total))
    w = np.asarray(pesos, float)
    sw = w.sum()
    if total <= 0 or sw <= 0:
        return np.zeros(len(w), int)
    exacto = total * w / sw
    base = np.floor(exacto).astype(int)
    faltan = total - base.sum()
    if faltan > 0:
        orden = np.argsort(-(exacto - base))
        base[orden[:faltan]] += 1
    return base


class Repartidor:
    """Envuelve al modelo para repartir filas de ingresantes por ciclo."""

    def __init__(self, ruta_par=PAR):
        if not os.path.exists(ruta_par):
            sys.exit("No encuentro %s. Ejecuta antes `python3 fase2.py`." % ruta_par)
        self.p = json.load(open(ruta_par, encoding="utf-8"))
        self.M = E.Modelo(self.p)
        self.ciclos = self.p["ciclosNuevos"]
        self.plan = self.p.get("planCiclos", {})
        self.plan_def = self.p.get("planDefecto", 10)
        self.apertura = self.p.get("sedeApertura", {})
        self.catalogo = set(self.p.get("carreras", []))
        # Apertura de los programas que no figuran en el histórico; la fija
        # `registrar_nuevos` a partir del propio archivo de entrada.
        self.nuevas = {}

    def registrar_nuevos(self, filas):
        """
        Localiza los programas sin historia y fija su semestre de apertura.

        Un programa que arranca no imparte todavía ciclos superiores: en su
        primer semestre sólo existe el ciclo 1, en el segundo el 1 y el 2, y
        así sucesivamente, a medida que avanza su primera cohorte. Es la misma
        maduración que la de una sede nueva. Sin esto, el reparto estimado
        colocaría convalidaciones en ciclos que el programa no imparte.

        El ciclo base es el mayor declarado en el semestre de apertura —por si
        el programa se lanza convalidando desde otro afín— y 1 si no se declara
        ninguno, que es el caso normal.
        """
        for carr, per, ciclo in filas:
            if carr in self.catalogo:
                continue
            ap = self.nuevas.setdefault(carr, {"inicio": per, "cicloBase": 1})
            if per < ap["inicio"]:
                ap["inicio"] = per
                ap["cicloBase"] = 1
            if ciclo and per == ap["inicio"]:
                ap["cicloBase"] = max(ap["cicloBase"], int(ciclo))
        return self.nuevas

    def tope(self, sede, carrera, periodo, plan_decl=0):
        """Ciclo máximo admisible: plan, maduración de sede y de programa."""
        t = plan_decl or self.plan.get(carrera, self.plan_def)
        ap = self.apertura.get(sede)
        if ap and ap.get("enMaduracion"):
            t = min(t, self.M.tope_sede(sede, periodo))
        nueva = self.nuevas.get(carrera)
        if nueva:
            avance = E.indice_periodo(periodo) - E.indice_periodo(nueva["inicio"])
            t = min(t, nueva["cicloBase"] + avance)
        return max(1, int(t))

    def reparto(self, sede, carrera, moda, periodo, plan_decl=0):
        parid = periodo % 100
        v, nef, niv = self.M.mezcla_ciclo(sede, carrera, moda, parid,
                                          tope=self.tope(sede, carrera, periodo, plan_decl))
        return v, nef, niv

    def fila(self, sede, carrera, moda, periodo, total, plan_decl=0):
        v, nef, niv = self.reparto(sede, carrera, moda, periodo, plan_decl)
        n = resto_mayor(total, v)
        return [(c, int(x)) for c, x in zip(self.ciclos, n) if x > 0], nef, niv


def hoja_reparto(wb, R, periodos=(1, 2)):
    """Cuadro auditable de las tasas que se han aplicado."""
    ws = wb.create_sheet("Tasas aplicadas")
    ws.append(["Sede", "Carrera", "Modalidad", "Paridad",
               "% ciclo 1", "% ciclo 2", "% ciclo 3", "% ciclo 4", "% ciclo 5+",
               "% que no entra en ciclo 1", "n efectivo", "Nivel de la cascada"])
    for j in range(1, 13):
        c = ws.cell(row=1, column=j)
        c.font = Font(bold=True)
        c.fill = CABF
        c.alignment = Alignment(wrap_text=True, vertical="center")
    combis = set()
    for k in R.p["nc_nivel_celda"]:
        s, c, m, pa = k.split("|")
        combis.add((s, c, m, int(pa)))
    for s, c, m, pa in sorted(combis):
        per = 2027 * 100 + pa
        v, nef, niv = R.M.mezcla_ciclo(s, c, m, pa, tope=R.tope(s, c, per))
        v5 = float(v[4:].sum())
        ws.append([s, c, m, "I" if pa == 1 else "II",
                   float(v[0]), float(v[1]), float(v[2]), float(v[3]), v5,
                   float(1 - v[0]), round(nef, 1), niv])
    for i in range(2, ws.max_row + 1):
        for j in range(5, 11):
            ws.cell(row=i, column=j).number_format = "0,00 %"
    for j, w in enumerate([13, 40, 16, 9, 10, 10, 10, 10, 10, 13, 11, 26], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:L%d" % ws.max_row
    return ws


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    solo_tabla = "--tabla" in sys.argv
    R = Repartidor()

    if solo_tabla or not args:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        hoja_reparto(wb, R)
        dest = os.path.join(AQUI, "..", "reparto_ingresantes_por_ciclo.xlsx")
        wb.save(dest)
        print("%s  ->  %d combinaciones" % (os.path.normpath(dest),
                                            wb["Tasas aplicadas"].max_row - 1))
        if not solo_tabla:
            print("\nUso: python3 repartir_ingresantes.py entrada.xlsx [salida.xlsx]")
        return

    entrada = args[0]
    salida = args[1] if len(args) > 1 else os.path.splitext(entrada)[0] + "_por_ciclo.xlsx"
    wb = openpyxl.load_workbook(entrada, data_only=True)
    ws = wb.active
    filas = list(ws.values)
    if not filas:
        sys.exit("El archivo está vacío.")
    cab = [str(x).strip().lower() if x is not None else "" for x in filas[0]]

    def col(*nombres):
        for n in nombres:
            for j, c in enumerate(cab):
                if c.startswith(n):
                    return j
        return None

    ip, isd, ic, im, inv = (col("periodo", "semestre"), col("sede"), col("carrera"),
                            col("modalidad"), col("nuevos", "ingresantes", "total"))
    ipl = col("ciclosplan", "ciclos plan", "plan")
    falta = [n for n, v in (("Periodo", ip), ("Sede", isd), ("Carrera", ic),
                            ("Modalidad", im), ("Nuevos", inv)) if v is None]
    if falta:
        sys.exit("Faltan columnas: %s\nCabecera leída: %s" % (", ".join(falta), cab))

    out = openpyxl.Workbook()
    o = out.active
    o.title = "Ingresantes"
    o.append(CAB)
    for j in range(1, len(CAB) + 1):
        o.cell(row=1, column=j).font = Font(bold=True)
        o.cell(row=1, column=j).fill = CABF

    # Primera pasada: fijar la apertura de los programas sin historia, porque
    # el tope de una fila depende del semestre en que arranca su programa.
    previas = []
    for f in filas[1:]:
        if f is None or all(x is None for x in f):
            continue
        try:
            previas.append((str(f[ic]).strip().upper(), leer_periodo(f[ip]), None))
        except (TypeError, ValueError):
            continue
    R.registrar_nuevos(previas)
    if R.nuevas:
        print("  programas sin historia: %s" % ", ".join(
            "%s (abre %s)" % (c, rotulo(a["inicio"])) for c, a in sorted(R.nuevas.items())))

    n_in = n_out = 0
    n_filas = 0
    avisos = []
    for i, f in enumerate(filas[1:], start=2):
        if f is None or all(x is None for x in f):
            continue
        try:
            per = leer_periodo(f[ip])
            sede = str(f[isd]).strip()
            carr = str(f[ic]).strip().upper()
            moda = str(f[im]).strip()
            tot = int(round(float(str(f[inv]).replace(",", ".").replace(" ", ""))))
        except (TypeError, ValueError):
            avisos.append("fila %d ilegible, descartada" % i)
            continue
        if tot <= 0:
            continue
        plan = 0
        if ipl is not None and f[ipl] not in (None, ""):
            try:
                plan = int(float(f[ipl]))
            except (TypeError, ValueError):
                plan = 0
        rep, nef, niv = R.fila(sede, carr, moda, per, tot, plan)
        n_in += tot
        n_filas += 1
        for c, n in rep:
            o.append([rotulo(per), sede, carr, moda, c, n, plan or None])
            n_out += n
        if nef < 30:
            avisos.append("fila %d (%s · %s · %s): poca evidencia propia, "
                          "se hereda del nivel «%s»" % (i, sede, carr[:22], moda, niv))

    for j, w in enumerate([11, 14, 42, 17, 8, 10, 12], start=1):
        o.column_dimensions[get_column_letter(j)].width = w
    o.freeze_panes = "A2"
    hoja_reparto(out, R)
    out.save(salida)

    print("%s" % os.path.normpath(salida))
    print("  %d filas de entrada  ->  %d filas por ciclo" % (n_filas, o.max_row - 1))
    print("  ingresantes declarados %s  ->  repartidos %s  %s"
          % ("{:,}".format(n_in).replace(",", " "),
             "{:,}".format(n_out).replace(",", " "),
             "CUADRA" if n_in == n_out else "!! DESCUADRE"))
    if avisos:
        print("  avisos (%d):" % len(avisos))
        for a in avisos[:8]:
            print("    - " + a)
        if len(avisos) > 8:
            print("    ... y %d más" % (len(avisos) - 8))


if __name__ == "__main__":
    main()
