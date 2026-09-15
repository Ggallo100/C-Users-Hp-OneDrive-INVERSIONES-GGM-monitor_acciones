#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera `../tablas_continuadores.xlsx`: la matrícula histórica clasificada por
condición de llegada —ingresante, regular, recuperado y reiniciado— a lo largo
de todos los semestres, abierta por carrera, por modalidad y por ciclo.

La clasificación es la misma que usa el modelo y se deriva del propio panel de
matrículas (ver `condicion_de` en estimar.py). No hay parámetros ni supuestos:
cada matrícula cae en una y sólo una condición, de modo que las cuatro suman
siempre el total del semestre.

    python3 tablas_condicion.py
"""
import os

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import estimar as E

AQUI = os.path.dirname(os.path.abspath(__file__))
SALIDA = os.path.join(AQUI, "..", "tablas_continuadores.xlsx")
COND = ["Ingresante", "Regular", "Recuperado", "Reiniciado"]

AZUL = "1F3864"
CAB = PatternFill("solid", fgColor="D9E2F3")
TIT = PatternFill("solid", fgColor=AZUL)
TOT = PatternFill("solid", fgColor="F2F2F2")
CENS = PatternFill("solid", fgColor="FBE4D5")
FINO = Side(style="thin", color="BFBFBF")
BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)


def rotulo(p):
    return "%d-%s" % (p // 100, "I" if p % 100 == 1 else "II")


def censurada(cond, p):
    """Semestres en que la condición no es observable por el inicio de la ventana."""
    if cond == "Recuperado":
        return p == 202301
    if cond == "Reiniciado":
        return p in (202301, 202302)
    return False


def bloque(ws, fila, titulo, tabla, periodos, cond=None, pct=False, cond_fila=False,
           etiq=None, tot_rotulo=None, tot_fn=None):
    """
    Escribe un bloque categoría × semestre y devuelve la fila siguiente.

    `cond` marca la condición cuando todo el bloque es de una sola —la censura
    afecta entonces a columnas enteras—; `cond_fila` la toma de la etiqueta de
    cada fila, que es el caso del cuadro institucional.
    """
    c = ws.cell(row=fila, column=1, value=titulo)
    c.font = Font(bold=True, color="FFFFFF", size=11)
    c.fill = TIT
    ws.merge_cells(start_row=fila, start_column=1,
                   end_row=fila, end_column=len(periodos) + 2)
    fila += 1
    ws.cell(row=fila, column=1, value=tabla.index.name or "Categoría").font = Font(bold=True)
    ws.cell(row=fila, column=1).fill = CAB
    ws.cell(row=fila, column=1).border = BORDE
    for j, p in enumerate(periodos):
        c = ws.cell(row=fila, column=2 + j, value=(etiq or rotulo)(p))
        c.font = Font(bold=True)
        c.fill = CENS if (cond and censurada(cond, p)) else CAB
        c.alignment = Alignment(horizontal="center")
        c.border = BORDE
    c = ws.cell(row=fila, column=2 + len(periodos),
                value=tot_rotulo or ("Total" if not pct else "Media sin censura"))
    c.font = Font(bold=True)
    c.fill = CAB
    c.alignment = Alignment(horizontal="center")
    c.border = BORDE
    fila += 1
    et_col = [(etiq or rotulo)(p) for p in periodos]
    pctcol = {j for j, e in enumerate(et_col) if str(e).strip().startswith("%")}
    for et, r in tabla.iterrows():
        es_tot = str(et).startswith("TOTAL")
        cf = str(et) if cond_fila else cond
        c = ws.cell(row=fila, column=1, value=str(et))
        c.border = BORDE
        if es_tot:
            c.font = Font(bold=True)
            c.fill = TOT
        for j, p in enumerate(periodos):
            v = r.get(p, 0)
            if j in pctcol:
                c = ws.cell(row=fila, column=2 + j, value=round(float(v), 2))
                c.number_format = '0.00" %"'
            else:
                c = ws.cell(row=fila, column=2 + j,
                            value=(float(v) / 100 if pct else int(v)))
                c.number_format = "0,0 %" if pct else "#,##0"
            c.border = BORDE
            if es_tot:
                c.font = Font(bold=True)
                c.fill = TOT
            elif cf and censurada(cf, p):
                c.fill = CENS
        if tot_fn == "ninguno":
            fila += 1
            continue
        if tot_fn == "suma" or (tot_fn is None and not pct):
            tot = r.sum()
        else:
            # La media omite los semestres en que la condición no es observable:
            # promediar ceros censurados subestimaría la proporción real.
            vis = [r.get(p, 0) for p in periodos if not (cf and censurada(cf, p))]
            tot = float(np.mean(vis)) if vis else 0.0
        c = ws.cell(row=fila, column=2 + len(periodos),
                    value=(float(tot) / 100 if pct else int(tot)))
        c.number_format = "0,0 %" if pct else "#,##0"
        c.font = Font(bold=True)
        c.fill = TOT
        c.border = BORDE
        fila += 1
    return fila + 1


def ancho(ws, primera=34, resto=11):
    ws.column_dimensions["A"].width = primera
    for j in range(2, 2 + 13):
        ws.column_dimensions[get_column_letter(j)].width = resto
    ws.freeze_panes = "B1"


def apertura(wb, nombre, b, campo, periodos, orden=None):
    """Una hoja con los cinco bloques de recuento y los tres de porcentaje."""
    ws = wb.create_sheet(nombre)
    ancho(ws, 34 if campo == "Carrera" else 20)
    fila = 1
    tot = pd.crosstab(b[campo], b["Periodo_real"])
    if orden is not None:
        tot = tot.reindex(orden).fillna(0)
    for cond in COND:
        t = pd.crosstab(b[b["condicion"] == cond][campo], b[b["condicion"] == cond]["Periodo_real"])
        t = t.reindex(index=tot.index, columns=periodos).fillna(0)
        t.index.name = campo
        t.loc["TOTAL institucional"] = t.sum()
        fila = bloque(ws, fila, "%s — matrículas" % cond.upper(), t, periodos, cond=cond)
    t = tot.reindex(columns=periodos).fillna(0)
    t.index.name = campo
    t.loc["TOTAL institucional"] = t.sum()
    fila = bloque(ws, fila, "TOTAL DE MATRÍCULA", t, periodos)
    for cond in ("Regular", "Recuperado", "Reiniciado"):
        num = pd.crosstab(b[b["condicion"] == cond][campo],
                          b[b["condicion"] == cond]["Periodo_real"]).reindex(
            index=tot.index, columns=periodos).fillna(0)
        den = tot.reindex(columns=periodos).fillna(0)
        pc = (100 * num / den.replace(0, np.nan)).fillna(0)
        pc.index.name = campo
        pc.loc["TOTAL institucional"] = (100 * num.sum() / den.sum()).fillna(0)
        fila = bloque(ws, fila, "%% DE %s SOBRE LA MATRÍCULA DE LA FILA" % cond.upper(),
                      pc, periodos, cond=cond, pct=True)
    return ws


def ingresantes(wb, b, periodos):
    """
    Dos hojas sobre el ciclo en que ENTRAN los ingresantes: no todos lo hacen
    en el primero, y el reparto tiene estructura.
    """
    ing = b[b["esNuevo"] == 1].copy()
    ing["Ciclo"] = ing["Ciclo"].astype(int)
    ciclos = sorted(ing["Ciclo"].unique())
    ident = lambda c: ("Ciclo %d" % c)

    # ---- hoja 1: institucional y forma del reparto -------------------------
    ws = wb.create_sheet("Ingresantes por ciclo")
    ancho(ws, 26, 11)
    t = pd.crosstab(ing["Periodo_real"], ing["Ciclo"]).reindex(
        index=periodos, columns=ciclos).fillna(0)
    t.index = [rotulo(p) for p in periodos]
    t.index.name = "Semestre"
    t.loc["TOTAL"] = t.sum()
    fila = bloque(ws, 1, "INGRESANTES POR CICLO DE INGRESO", t, ciclos, etiq=ident)
    pc = 100 * t.iloc[:-1].div(t.iloc[:-1].sum(axis=1), axis=0)
    pc.index.name = "Semestre"
    pc.loc["TOTAL"] = 100 * t.iloc[:-1].sum() / t.iloc[:-1].sum().sum()
    fila = bloque(ws, fila, "REPARTO DE CADA SEMESTRE (%)", pc, ciclos,
                  pct=True, etiq=ident, tot_rotulo="Suma", tot_fn="suma")

    # nivel frente a proporción: el recuento de ingreso alto es mucho más
    # estable que su cuota, porque la que se mueve es la campaña de ciclo 1
    n1 = t.iloc[:-1][1]
    nm = t.iloc[:-1].drop(columns=[1]).sum(axis=1)
    d = pd.DataFrame({"Ingresantes en ciclo 1": n1,
                      "Ingresantes por encima del ciclo 1": nm}).T
    d.index.name = "Recuento"
    fila = bloque(ws, fila, "NIVEL FRENTE A PROPORCIÓN: LA CLAVE DEL PATRÓN",
                  d, list(d.columns), etiq=lambda x: x, tot_rotulo="Total")
    q = pd.DataFrame({"% que NO entra en ciclo 1": 100 * nm / (n1 + nm)}).T
    q.index.name = "Proporción"
    fila = bloque(ws, fila, "", q, list(q.columns), pct=True,
                  etiq=lambda x: x, tot_rotulo="Media")

    # forma condicional por semestre
    alto = ing[ing["Ciclo"] > 1]
    cl = sorted(alto["Ciclo"].clip(upper=5).unique())
    t2 = pd.crosstab(alto["Periodo_real"], alto["Ciclo"].clip(upper=5)).reindex(
        index=periodos, columns=cl).fillna(0)
    pc2 = 100 * t2.div(t2.sum(axis=1), axis=0)
    pc2.index = [rotulo(p) for p in periodos]
    pc2.index.name = "Semestre"
    pc2.loc["TOTAL"] = 100 * t2.sum() / t2.sum().sum()
    fila = bloque(ws, fila, "DE LOS QUE ENTRAN POR ENCIMA DEL CICLO 1, ¿A QUÉ CICLO LLEGAN? (%)",
                  pc2, cl, pct=True,
                  etiq=lambda c: ("Ciclo 5 o más" if c == 5 else "Ciclo %d" % c),
                  tot_rotulo="Suma", tot_fn="suma")

    # ---- hoja 2: carrera y modalidad ---------------------------------------
    ws = wb.create_sheet("Ingr. carrera y modalidad")
    ancho(ws, 44, 11)
    fila = 1
    for campo, titulo in (("Carrera", "CARRERA"), ("Modalidad", "MODALIDAD"), ("Sede", "SEDE")):
        t = pd.crosstab(ing[campo], ing["Ciclo"]).reindex(columns=ciclos).fillna(0)
        t.index.name = campo
        t["Total"] = t.sum(axis=1)
        t["No entra en ciclo 1"] = t["Total"] - t[1]
        t["% que no entra en ciclo 1"] = (100 * t["No entra en ciclo 1"] / t["Total"]).round(2)
        t = t.sort_values("% que no entra en ciclo 1", ascending=False)
        t.loc["TOTAL institucional"] = t.sum()
        t.loc["TOTAL institucional", "% que no entra en ciclo 1"] = round(
            100 * t.loc["TOTAL institucional", "No entra en ciclo 1"]
            / t.loc["TOTAL institucional", "Total"], 2)
        cols = ciclos + ["Total", "No entra en ciclo 1", "% que no entra en ciclo 1"]
        fila = bloque(ws, fila, "INGRESANTES POR %s Y CICLO DE INGRESO" % titulo,
                      t[cols], cols,
                      etiq=lambda c: (("Ciclo %d" % c) if isinstance(c, (int, np.integer)) else c),
                      tot_fn="ninguno")
    # cruce carrera x modalidad del nivel
    t = ing.pivot_table(index="Carrera", columns="Modalidad", values="Ciclo",
                        aggfunc=lambda x: 100 * (np.asarray(x) > 1).mean()).round(1)
    n = ing.pivot_table(index="Carrera", columns="Modalidad", values="Ciclo", aggfunc="size")
    t = t.where(n >= 50)
    t.index.name = "Carrera"
    bloque(ws, fila, "% QUE NO ENTRA EN CICLO 1, POR CARRERA Y MODALIDAD (celdas con n ≥ 50)",
           t.fillna(0), list(t.columns), pct=True,
           etiq=lambda x: x, tot_rotulo="Media")
    return ws


def leeme(wb, b, periodos, n):
    ws = wb.create_sheet("Léeme", 0)
    ws.column_dimensions["A"].width = 120
    lineas = [
        ("Clasificación de la matrícula por continuidad", True),
        ("", False),
        ("Fuente: %s" % os.path.basename(E.RUTA), False),
        ("Matrículas analizadas: %s registros, %d semestres (%s a %s)."
         % ("{:,}".format(n).replace(",", " "), len(periodos),
            rotulo(periodos[0]), rotulo(periodos[-1])), False),
        ("Unidad de análisis: una fila por estudiante y semestre.", False),
        ("", False),
        ("Las cuatro condiciones", True),
        ("INGRESANTE   Primera matrícula del estudiante (campo Nuevos de la base).", False),
        ("REGULAR      Se matriculó también el semestre inmediato anterior y cambia de ciclo:", False),
        ("             continúa sin interrupción.", False),
        ("RECUPERADO   Se matriculó también el semestre inmediato anterior, pero vuelve AL MISMO", False),
        ("             CICLO: perdió el que cursaba —figura como desertor de ese ciclo— y lo", False),
        ("             retoma al semestre siguiente.", False),
        ("REINICIADO   Interrumpió uno o más semestres y volvió a matricularse.", False),
        ("", False),
        ("Las cuatro condiciones parten la matrícula sin solaparse: cada matrícula cae en una y", False),
        ("sólo una, de modo que suman exactamente el total del semestre. No se declaran en ningún", False),
        ("sitio ni se estiman: se derivan del propio panel, comparando cada matrícula con la", False),
        ("anterior del mismo estudiante (la brecha en semestres y el salto de ciclo).", False),
        ("", False),
        ("Por qué no se usa la columna «Desertor» de la base", True),
        ("Su regla interna es «no se matriculó el semestre inmediato siguiente» —se cumple con cero", False),
        ("excepciones en los seis semestres cerrados— de modo que por construcción excluye a quien", False),
        ("sí se matriculó, que es justamente el recuperado: definirlo con esa marca daría una", False),
        ("categoría vacía. Además está congelada a una fecha anterior a la campaña de 2026-II, y", False),
        ("marca como desertores al 65,5 % de los matriculados de 2026-I y al 0 % de los de 2026-II.", False),
        ("", False),
        ("Censura al inicio de la ventana  (celdas en naranja)", True),
        ("Para ver a un recuperado hace falta un semestre previo y para ver a un reiniciado, dos.", False),
        ("2023-I no puede aportar ninguna de las dos categorías y 2023-II no puede aportar", False),
        ("reiniciados: esas matrículas se cuentan como regulares. No es que no existieran, es que", False),
        ("no son observables. A partir de 2024-I las cifras son comparables entre sí.", False),
        ("", False),
        ("Hojas de este libro", True),
        ("Institucional   Los totales de la universidad, semestre a semestre.", False),
        ("Por carrera     Las 25 carreras, con los cinco bloques de recuento y los tres de %.", False),
        ("Por modalidad   Presencial, semipresencial y a distancia.", False),
        ("Por ciclo       Los once ciclos del plan.", False),
        ("Ingresantes por ciclo        En qué ciclo entran los nuevos: el 96,2 % en el primero,", False),
        ("                             pero 1 205 no. Reparto por semestre y forma del ingreso alto.", False),
        ("Ingr. carrera y modalidad    El mismo reparto abierto por carrera, modalidad y sede.", False),
        ("Detalle         Formato largo (una fila por combinación) para tablas dinámicas.", False),
        ("                Incluye la sede, por si se quiere abrir también por ella.", False),
        ("", False),
        ("Los porcentajes se calculan siempre sobre la matrícula total de esa fila y semestre.", False),
    ]
    for i, (t, neg) in enumerate(lineas, start=1):
        c = ws.cell(row=i, column=1, value=t)
        if neg:
            c.font = Font(bold=True, size=12, color=AZUL)
    return ws


def main():
    df = E.cargar()
    b, per, idx = E.panel(df)
    b = b.rename(columns={"Modalidad_estudios": "Modalidad"})
    periodos = list(per)

    wb = Workbook()
    wb.remove(wb.active)
    leeme(wb, b, periodos, len(b))

    # ---- institucional -----------------------------------------------------
    ws = wb.create_sheet("Institucional")
    ancho(ws, 24)
    t = pd.crosstab(b["condicion"], b["Periodo_real"]).reindex(
        index=COND, columns=periodos).fillna(0)
    t.index.name = "Condición"
    t.loc["TOTAL matrícula"] = t.sum()
    fila = bloque(ws, 1, "MATRÍCULA POR CONDICIÓN DE LLEGADA", t, periodos)
    pc = 100 * t.iloc[:-1] / t.iloc[:-1].sum()
    pc.index.name = "Condición"
    pc.loc["TOTAL"] = pc.sum()
    bloque(ws, fila, "COMPOSICIÓN DE CADA SEMESTRE (%)", pc, periodos,
           pct=True, cond_fila=True)

    # ---- aperturas ---------------------------------------------------------
    apertura(wb, "Por carrera", b, "Carrera", periodos)
    apertura(wb, "Por modalidad", b, "Modalidad", periodos)
    b["Ciclo"] = b["Ciclo"].astype(int)
    apertura(wb, "Por ciclo", b, "Ciclo", periodos, orden=sorted(b["Ciclo"].unique()))

    ingresantes(wb, b, periodos)

    # ---- detalle en formato largo ------------------------------------------
    ws = wb.create_sheet("Detalle")
    d = (b.groupby(["Periodo_real", "Sede", "Carrera", "Modalidad", "Ciclo", "condicion"])
         .size().reset_index(name="Matriculados"))
    d.insert(0, "Semestre", d["Periodo_real"].map(rotulo))
    d = d.drop(columns=["Periodo_real"]).rename(columns={"condicion": "Condición"})
    cabs = list(d.columns)
    for j, h in enumerate(cabs, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = CAB
        c.border = BORDE
    for i, r in enumerate(d.itertuples(index=False), start=2):
        for j, v in enumerate(r, start=1):
            ws.cell(row=i, column=j, value=v)
    for j, h in enumerate(cabs, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(38, len(h) + 22))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cabs)), len(d) + 1)

    wb.save(SALIDA)
    print("%s  ->  %d hojas, %d filas de detalle"
          % (os.path.normpath(SALIDA), len(wb.sheetnames), len(d)))
    print("total clasificado: %s matrículas" % "{:,}".format(len(b)).replace(",", " "))


if __name__ == "__main__":
    main()
